import csv
import io
from typing import Tuple
from django.http import HttpResponse
from django.utils import timezone
from accounts.models import StudentProfile, ResearchGroup
from evaluations.models import StudentAssessment, AssessmentType
from chatbot.models import ChatSession, ChatMessage, MessageRole
from research.models import ExportAuditLog, ExportDatasetType, ExportFormat

def export_participants_data(file_format: str = 'csv', researcher_user=None) -> HttpResponse:
    """
    Exports the participants registry dataset.
    """
    students = StudentProfile.objects.all().order_by('student_code')
    records_count = students.count()

    headers = [
        'student_code', 'group_code', 'group_name', 'current_stage',
        'consent_given', 'consent_timestamp', 'intervention_completed',
        'intervention_completed_at', 'is_active', 'registered_at'
    ]

    rows = []
    for s in students:
        rows.append([
            s.student_code,
            s.group,
            s.get_group_display(),
            s.get_stage_display(),
            '1' if s.consent_given else '0',
            s.consent_timestamp.strftime('%Y-%m-%d %H:%M:%S') if s.consent_timestamp else '',
            '1' if s.intervention_completed else '0',
            s.intervention_completed_at.strftime('%Y-%m-%d %H:%M:%S') if s.intervention_completed_at else '',
            '1' if s.is_active else '0',
            s.created_at.strftime('%Y-%m-%d %H:%M:%S') if s.created_at else '',
        ])

    # Log audit
    ExportAuditLog.objects.create(
        researcher=researcher_user,
        dataset_type=ExportDatasetType.PARTICIPANTS,
        file_format=ExportFormat.CSV if file_format == 'csv' else ExportFormat.XLSX,
        records_count=records_count,
        exported_at=timezone.now()
    )

    filename = f"previene_participantes_{timezone.now().strftime('%Y%m%d_%H%M%S')}"
    return _build_file_response(headers, rows, filename, file_format)

def export_evaluations_data(file_format: str = 'csv', researcher_user=None) -> HttpResponse:
    """
    Exports paired Pretest and Postest evaluations dataset for SPSS, R, Jamovi or Python.
    Calculates dimensional subscores and deltas automatically.
    """
    students = StudentProfile.objects.prefetch_related('student_assessments').all().order_by('student_code')
    records_count = students.count()

    headers = [
        'student_code', 'group_code', 'group_name',
        'pretest_completed', 'pretest_dim1_knowledge', 'pretest_dim2_practices', 'pretest_total_score', 'pretest_date',
        'postest_completed', 'postest_dim1_knowledge', 'postest_dim2_practices', 'postest_total_score', 'postest_date',
        'delta_dim1', 'delta_dim2', 'delta_total'
    ]

    rows = []
    for s in students:
        pre = s.student_assessments.filter(assessment_type=AssessmentType.PRETEST, is_submitted=True).first()
        post = s.student_assessments.filter(assessment_type=AssessmentType.POSTEST, is_submitted=True).first()

        pre_done = '1' if pre else '0'
        pre_d1 = round(pre.score_dimension_1, 2) if pre else ''
        pre_d2 = round(pre.score_dimension_2, 2) if pre else ''
        pre_tot = round(pre.total_score, 2) if pre else ''
        pre_date = pre.completed_at.strftime('%Y-%m-%d %H:%M:%S') if (pre and pre.completed_at) else ''

        post_done = '1' if post else '0'
        post_d1 = round(post.score_dimension_1, 2) if post else ''
        post_d2 = round(post.score_dimension_2, 2) if post else ''
        post_tot = round(post.total_score, 2) if post else ''
        post_date = post.completed_at.strftime('%Y-%m-%d %H:%M:%S') if (post and post.completed_at) else ''

        # Calculate Deltas (Post - Pre)
        if pre and post:
            delta_d1 = round(post.score_dimension_1 - pre.score_dimension_1, 2)
            delta_d2 = round(post.score_dimension_2 - pre.score_dimension_2, 2)
            delta_tot = round(post.total_score - pre.total_score, 2)
        else:
            delta_d1 = ''
            delta_d2 = ''
            delta_tot = ''

        rows.append([
            s.student_code,
            s.group,
            s.get_group_display(),
            pre_done,
            pre_d1,
            pre_d2,
            pre_tot,
            pre_date,
            post_done,
            post_d1,
            post_d2,
            post_tot,
            post_date,
            delta_d1,
            delta_d2,
            delta_tot
        ])

    ExportAuditLog.objects.create(
        researcher=researcher_user,
        dataset_type=ExportDatasetType.EVALUATIONS,
        file_format=ExportFormat.CSV if file_format == 'csv' else ExportFormat.XLSX,
        records_count=records_count,
        exported_at=timezone.now()
    )

    filename = f"previene_evaluaciones_{timezone.now().strftime('%Y%m%d_%H%M%S')}"
    return _build_file_response(headers, rows, filename, file_format)

def export_chatbot_usage_data(file_format: str = 'csv', researcher_user=None) -> HttpResponse:
    """
    Exports the Previene+ chatbot interaction metrics dataset.
    """
    sessions = ChatSession.objects.select_related('student').prefetch_related('messages').all().order_by('-started_at')
    records_count = sessions.count()

    headers = [
        'student_code', 'group_code', 'session_uuid', 'total_messages',
        'user_messages_count', 'assistant_messages_count', 'duration_seconds',
        'duration_minutes', 'topics_consulted', 'started_at', 'ended_at', 'is_active'
    ]

    rows = []
    for sess in sessions:
        user_msgs = sess.messages.filter(role=MessageRole.USER).count()
        asst_msgs = sess.messages.filter(role=MessageRole.ASSISTANT).count()
        topics = set()
        for msg in sess.messages.all():
            if msg.topics_detected:
                for t in msg.topics_detected.split(','):
                    if t.strip():
                        topics.add(t.strip())

        dur_sec = sess.duration_seconds
        dur_min = round(dur_sec / 60.0, 2)

        rows.append([
            sess.student.student_code,
            sess.student.group,
            str(sess.session_uuid),
            sess.total_messages,
            user_msgs,
            asst_msgs,
            dur_sec,
            dur_min,
            "; ".join(sorted(topics)),
            sess.started_at.strftime('%Y-%m-%d %H:%M:%S') if sess.started_at else '',
            sess.ended_at.strftime('%Y-%m-%d %H:%M:%S') if sess.ended_at else '',
            '1' if sess.is_active else '0',
        ])

    ExportAuditLog.objects.create(
        researcher=researcher_user,
        dataset_type=ExportDatasetType.CHATBOT_USAGE,
        file_format=ExportFormat.CSV if file_format == 'csv' else ExportFormat.XLSX,
        records_count=records_count,
        exported_at=timezone.now()
    )

    filename = f"previene_uso_chatbot_{timezone.now().strftime('%Y%m%d_%H%M%S')}"
    return _build_file_response(headers, rows, filename, file_format)

def _build_file_response(headers, rows, filename: str, file_format: str) -> HttpResponse:
    """Builds CSV or XLSX downloadable HTTP response."""
    if file_format == 'xlsx':
        try:
            import openpyxl
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Previene+ Data"

            # Header style
            ws.append(headers)
            for cell in ws[1]:
                cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
                cell.fill = openpyxl.styles.PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")

            for r in rows:
                ws.append(r)

            output = io.BytesIO()
            wb.save(output)
            output.seek(0)

            response = HttpResponse(
                output.getvalue(),
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            )
            response['Content-Disposition'] = f'attachment; filename="{filename}.xlsx"'
            return response
        except ImportError:
            pass

    # Default CSV output with UTF-8 BOM for flawless Excel encoding
    output = io.StringIO()
    writer = csv.writer(output, delimiter=',', quoting=csv.QUOTE_MINIMAL)
    writer.writerow(headers)
    for r in rows:
        writer.writerow(r)

    csv_data = output.getvalue()
    response = HttpResponse(
        '\ufeff' + csv_data,
        content_type='text/csv; charset=utf-8'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}.csv"'
    return response
